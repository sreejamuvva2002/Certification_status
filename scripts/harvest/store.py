"""JSONL/CSV artifact I/O. Every stage reads one artifact and writes another.

Stages are resumable by reading the id set already present in their own output,
the same way `load_done()` works in check_certifications.py. Writes are
flushed per record so a killed run keeps everything it had finished.
"""
from __future__ import annotations

import csv
import json
import os
import sys
import threading

_WRITE_LOCK = threading.Lock()


def ensure_dir(path: str) -> None:
    os.makedirs(path, exist_ok=True)


def read_jsonl(path: str) -> list[dict]:
    if not os.path.exists(path):
        return []
    out = []
    with open(path, "r", encoding="utf-8") as fh:
        for line_no, line in enumerate(fh, 1):
            line = line.strip()
            if not line:
                continue
            try:
                out.append(json.loads(line))
            except Exception as e:
                print(f"  ! {os.path.basename(path)}:{line_no} unparseable, skipped ({e})",
                      file=sys.stderr)
    return out


def append_jsonl(path: str, record: dict) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with _WRITE_LOCK, open(path, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")
        fh.flush()


def write_jsonl(path: str, records: list[dict]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        for r in records:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    os.replace(tmp, path)


def done_ids(path: str, key: str) -> set[str]:
    """Ids already present in a stage's own output — the resume set."""
    return {r.get(key) for r in read_jsonl(path) if r.get(key)}


def truncate(path: str) -> None:
    if os.path.exists(path):
        os.remove(path)


def write_csv(path: str, rows: list[dict], columns: list[str]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in columns})
    os.replace(tmp, path)


def read_csv(path: str) -> list[dict]:
    if not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def md_cell(value: object) -> str:
    """Markdown table cell: escape pipes, flatten newlines, em-dash for empty."""
    s = "" if value is None else str(value)
    s = s.replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()
    s = " ".join(s.split())
    return s or "—"


def md_table(rows: list[dict], columns: list[tuple[str, str]]) -> str:
    head = "| " + " | ".join(label for _, label in columns) + " |"
    rule = "| " + " | ".join("---" for _ in columns) + " |"
    body = [
        "| " + " | ".join(md_cell(r.get(key, "")) for key, _ in columns) + " |"
        for r in rows
    ]
    return "\n".join([head, rule] + body)


def write_xlsx(path: str, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    """Write a multi-sheet workbook. sheets = {name: (columns, rows)}."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        print("  ! openpyxl not installed — skipping xlsx", file=sys.stderr)
        return
    wb = Workbook()
    wb.remove(wb.active)
    for name, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
        ws.freeze_panes = "A2"
        for i, col in enumerate(columns, 1):
            width = max(12, min(48, len(col) + 4))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
