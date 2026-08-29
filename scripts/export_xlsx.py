"""Export the gap-fill result as Certification_Status_v2.xlsx.

Sheet 1 keeps the exact 7-column shape of the user's own outputs/Certification_Status.xlsx
(Record No. | Company | Likely Georgia facility/location | Certification / Standard | Scope |
Source type | Source URL) so it drops into the same downstream use, but each row is now a
canonical standard rather than a raw string, de-duplicated per company.

Sheet 2 carries what the 7-column shape cannot: the certificate identity pass C recovered
(registrar, number, issue, expiry) and — the column that changes every headline — the
facility scope the certificate was evidenced at.

Sheet 3 is the needs-review queue, so nothing that was held back from the counts is invisible.

The original xlsx is never touched; it is in run1_manifest.sha256 and stays byte-identical.

Usage:  python scripts/export_xlsx.py [--dir outputs/run2_gapfill]
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

PKG_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DIR = PKG_ROOT / "outputs" / "run2_gapfill"
NONE_IDS = {"__none__", "__no_ev__", "__no_details__"}

SHEET1 = ["Record No.", "Company", "Likely Georgia facility/location",
          "Certification / Standard", "Scope", "Source type", "Source URL"]
SHEET2 = ["Record No.", "Company", "Certification / Standard", "Family", "Domain",
          "Facility scope", "Certification status", "Certification body",
          "Certificate number", "Issue date", "Expiry date", "Scope", "Source type",
          "Source URL", "Evidence quote", "Confidence", "Found in run", "Notes"]
SHEET3 = ["Record No.", "Company", "Certification", "Reason", "Pass", "Evidence quote",
          "Source URL", "Notes"]


def _read(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _style(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def build(out_dir: Path) -> Path:
    rows = _read(out_dir / "certifications_normalized.csv")
    if not rows:
        raise SystemExit(f"no certifications_normalized.csv in {out_dir} — run normalize_certs.py")
    winners = [r for r in rows if r["is_winner"] == "1"]

    def sort_key(r: dict) -> tuple:
        try:
            return (int(r["record_no"]), r["canonical_id"])
        except ValueError:
            return (10 ** 9, r["canonical_id"])

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Certification_Status_v2"
    ws.append(SHEET1)
    for r in sorted(winners, key=sort_key):
        cert = ("none identified" if r["canonical_id"] in NONE_IDS else r["canonical_id"])
        ws.append([r["record_no"], r["company"], r["location"], cert,
                   r["scope"] or "—", r["source_type"] or "—", r["source_url"] or "—"])
    _style(ws, [10, 34, 28, 26, 52, 18, 46])

    ws2 = wb.create_sheet("Detail_and_scope")
    ws2.append(SHEET2)
    for r in sorted([x for x in winners if x["canonical_id"] not in NONE_IDS], key=sort_key):
        ws2.append([r["record_no"], r["company"], r["canonical_id"], r["family"], r["domain"],
                    r["scope_bucket"], r["cert_status"], r["certification_body"],
                    r["reference_no"], r["issue_date"], r["expiry_date"], r["scope"],
                    r["source_type"], r["source_url"], r["evidence_quote"], r["confidence"],
                    r["source_run"], r["notes"]])
    _style(ws2, [10, 30, 22, 22, 20, 20, 20, 18, 20, 14, 16, 44, 16, 40, 52, 12, 12, 40])

    ws3 = wb.create_sheet("Needs_review")
    ws3.append(SHEET3)
    for r in _read(out_dir / "needs_review.csv"):
        ws3.append([r["record_no"], r["company"], r["certification"], r["reason"], r["pass"],
                    r["evidence_quote"], r["source_url"], r["notes"]])
    _style(ws3, [10, 30, 24, 28, 8, 52, 40, 36])

    path = out_dir / "Certification_Status_v2.xlsx"
    wb.save(path)
    print(f"wrote {path}: {ws.max_row - 1} status rows, {ws2.max_row - 1} detail rows, "
          f"{ws3.max_row - 1} review rows")
    return path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", type=Path, default=DEFAULT_DIR)
    build(ap.parse_args().dir)


if __name__ == "__main__":
    main()
